from __future__ import annotations

import hashlib
import json
import os
import re
import shutil
import signal
import sqlite3
import subprocess
import sys
import threading
import time
import uuid
import csv
from zipfile import BadZipFile
from collections import deque
from pathlib import Path
from typing import Any, Optional


class CliProcessSupervisor:
    """Presentation-layer adapter around the canonical process_all_pos.py pipeline."""

    @staticmethod
    def _process_creationflags() -> int:
        if os.name != "nt":
            return 0
        return getattr(subprocess, "CREATE_NEW_PROCESS_GROUP", 0) | getattr(
            subprocess, "CREATE_NO_WINDOW", 0x08000000
        )

    def __init__(self) -> None:
        self.project_root = Path(__file__).resolve().parents[2]
        self.script = self.project_root / "process_all_pos.py"
        self.db_path = Path.home() / ".contract_downloader" / "cli_sessions.db"
        self.default_download_root = Path.home() / "Downloads" / "CoupaAttachments"
        self.process: Optional[subprocess.Popen[str]] = None
        self.input_path: Optional[str] = None
        self.download_root: Optional[str] = None
        self.run_dir: Optional[str] = None
        self.session_id: Optional[int] = None
        self.source_session_id: Optional[int] = None
        self.previous_max_session_id = 0
        self.started_at = 0.0
        self.stop_requested = False
        self.cancel_requested = False
        self._paused_for_auth = False
        self._logs: deque[dict[str, str]] = deque(maxlen=100)
        self._progress_samples: deque[tuple[float, int]] = deque(maxlen=180)
        self._last_processed = 0
        self._last_progress_at = 0.0
        self._lock = threading.RLock()
        self._metadata_path = Path.home() / ".contract_downloader" / "gui_cli_sessions.json"
        self._session_metadata: dict[str, dict[str, str]] = self._load_metadata()
        self._schema_ready = False

    def _load_metadata(self) -> dict[str, dict[str, str]]:
        try:
            data = json.loads(self._metadata_path.read_text(encoding="utf-8"))
            return data if isinstance(data, dict) else {}
        except (OSError, ValueError):
            return {}

    def _remember_session(self, session_id: int) -> None:
        if not self.input_path or not session_id:
            return
        value = {"input_path": self.input_path, "run_dir": self.run_dir or ""}
        if self._session_metadata.get(str(session_id)) == value:
            return
        self._session_metadata[str(session_id)] = value
        try:
            self._metadata_path.parent.mkdir(parents=True, exist_ok=True)
            self._metadata_path.write_text(json.dumps(self._session_metadata, indent=2), encoding="utf-8")
        except OSError:
            pass

    def _connect(self) -> sqlite3.Connection:
        self.db_path.parent.mkdir(parents=True, exist_ok=True)
        conn = sqlite3.connect(str(self.db_path), timeout=30)
        conn.row_factory = sqlite3.Row
        if not self._schema_ready:
            for column, definition in {
                "concurrency": "INTEGER DEFAULT 11",
                "duration_seconds": "REAL",
                "input_file_path": "TEXT",
                "input_file_blob": "BLOB",
                "input_file_sha256": "TEXT",
                "input_file_size": "INTEGER",
                "description": "TEXT",
            }.items():
                try:
                    conn.execute(f"ALTER TABLE sessions ADD COLUMN {column} {definition}")
                except sqlite3.OperationalError:
                    pass
            conn.execute(
                """
                CREATE TABLE IF NOT EXISTS retry_events (
                    id INTEGER PRIMARY KEY AUTOINCREMENT,
                    session_id INTEGER NOT NULL,
                    po_number TEXT,
                    requested_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
                    completed_at TIMESTAMP,
                    status_before TEXT,
                    status_after TEXT,
                    error_message TEXT
                )
                """
            )
            try:
                conn.execute("ALTER TABLE po_downloads ADD COLUMN remarks TEXT")
            except sqlite3.OperationalError:
                pass
            conn.execute(
                """
                CREATE TABLE IF NOT EXISTS retry_attempts (
                    id INTEGER PRIMARY KEY AUTOINCREMENT,
                    session_id INTEGER NOT NULL,
                    original_po_number TEXT NOT NULL,
                    edited_po_number TEXT NOT NULL,
                    staging_dir TEXT NOT NULL,
                    status TEXT DEFAULT 'RUNNING',
                    error_message TEXT,
                    requested_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
                    completed_at TIMESTAMP
                )
                """
            )
            conn.commit()
            self._schema_ready = True
        return conn

    def _latest_session_id(self) -> Optional[int]:
        try:
            with self._connect() as conn:
                row = conn.execute(
                    "SELECT id FROM sessions WHERE id > ? ORDER BY id DESC LIMIT 1",
                    (self.previous_max_session_id,),
                ).fetchone()
                return int(row["id"]) if row else None
        except sqlite3.Error:
            return None

    @staticmethod
    def _actionable_log_entry(line: str) -> Optional[dict[str, str]]:
        """Convert verbose CLI output into concise, user-facing run events."""
        if not line or line.startswith("[TIMING]") or set(line) == {"="}:
            return None
        if "[AUTH] Chaves:" in line:
            return None
        if line.startswith("[ERROR][PO "):
            return {"type": "Error", "message": f"Error: {line[len('[ERROR]'):].lstrip()}"}

        progress = re.search(
            r"\[\s*(\d+)/(\d+)\]\s+ok=(\d+)\s+err=(\d+)\s+files=(\d+)",
            line,
        )
        if progress:
            done, total, success, errors, files = progress.groups()
            log_type = "Warning" if int(errors) else "Info"
            return {
                "type": log_type,
                "message": (
                    f"Progress: {done}/{total} POs · {success} succeeded · "
                    f"{errors} failed · {files} attachments"
                ),
            }

        for marker in ("Run folder:", "Pasta desta execucao:"):
            if marker in line:
                return {"type": "System", "message": f"Run folder: {line.split(marker, 1)[1].strip()}"}
        if line.startswith("Downloads:") or line.startswith("Downloads: "):
            return {"type": "System", "message": line}
        if "Cookies carregados do cache" in line or "Cached cookies loaded" in line:
            return {"type": "System", "message": "Checking the cached Coupa session…"}
        if "Cookies validos" in line or "Coupa session is valid" in line:
            return {"type": "Success", "message": "Coupa session validated."}
        if "Cookies expirados" in line or "Cached session expired" in line:
            return {"type": "Warning", "message": "The cached Coupa session expired; sign-in is required."}
        if "[AUTH][PAUSED]" in line or "[RUN_PAUSED]" in line:
            return {
                "type": "Warning",
                "message": "Run paused for authentication. Close Edge, sign in if requested, then resume the pending POs.",
            }
        if "perfil Work do Edge" in line or "perfil existente do Edge" in line:
            return {"type": "System", "message": "Opening Edge with the existing user profile."}
        if "perfil persistente do app" in line or "perfil persistente do aplicativo" in line:
            return {"type": "System", "message": "Opening Edge with the separate Coupa sign-in profile."}
        if "Aguardando login" in line:
            return {"type": "System", "message": "Waiting for Coupa sign-in in Edge…"}
        if "Login detectado" in line:
            return {"type": "Success", "message": "Coupa sign-in detected; validating the session…"}
        if "cookies capturados" in line.lower() or "Cookies extraidos" in line:
            return {"type": "Success", "message": "Coupa session captured securely."}

        imported = re.search(r"\[INFO\]\s+(\d+) POs (?:importadas|imported)", line)
        if imported:
            return {"type": "System", "message": f"{imported.group(1)} purchase orders queued."}
        workers = re.search(r"(?:Processando|Processing) (\d+) POs (?:com|with) (\d+)(?: concurrent)? workers", line)
        if workers:
            return {
                "type": "System",
                "message": f"Processing {workers.group(1)} POs with {workers.group(2)} concurrent workers.",
            }
        if line.startswith("[INFO] Importando "):
            return {"type": "System", "message": f"Reading input: {line.split('Importando ', 1)[1].rstrip('.')}"}
        if line.startswith("[INFO] Reading input: "):
            return {"type": "System", "message": line.removeprefix("[INFO] ")}
        if "Retry in-place" in line or "Retry individual" in line or "Retry incomplete" in line:
            return {"type": "System", "message": line.removeprefix("[INFO] ")}

        result = re.search(r"Success:\s*(\d+)\s*\|\s*Failed:\s*(\d+)", line)
        if result:
            log_type = "Warning" if int(result.group(2)) else "Success"
            return {
                "type": log_type,
                "message": f"Run result: {result.group(1)} succeeded · {result.group(2)} failed.",
            }
        attachments = re.search(r"Total (?:de anexos|attachments):\s*(\d+)", line)
        if attachments:
            return {"type": "Info", "message": f"Attachments found: {attachments.group(1)}."}
        if line.startswith("[MSG2PDF]"):
            if "Nenhum arquivo" in line or "No .msg files found" in line:
                return None
            message = re.sub(r"^\[MSG2PDF\](?:\[(?:WARN|WARNING)\])?\s*", "Email conversion: ", line)
            return {"type": "Warning" if "WARN" in line else "Info", "message": message}
        if line.startswith("[DEDUP]"):
            message = re.sub(r"^\[DEDUP\](?:\[(?:WARN|WARNING)\])?\s*", "Deduplication: ", line)
            return {"type": "Warning" if "WARN" in line else "Info", "message": message}
        if line.startswith("[REPORT]"):
            is_error = "[ERROR]" in line or "[ERRO]" in line
            message = line.replace("[REPORT][ERROR]", "Report:").replace("[REPORT][ERRO]", "Report:").replace("[REPORT]", "Report:")
            return {"type": "Error" if is_error else "Success", "message": message}
        if "rate limits (429)" in line or "Coupa rate limits (429)" in line:
            return {"type": "Warning", "message": line.replace("[AVISO]", "").replace("[WARNING]", "").strip()}
        if "falhas de autenticacao" in line or "Authentication failures" in line:
            return {"type": "Warning", "message": line.replace("[AVISO]", "").replace("[WARNING]", "").strip()}
        if "[ERRO" in line or "[ERROR" in line or "traceback" in line.lower() or "exception" in line.lower() or re.search(r"\w+Error:", line):
            return {"type": "Error", "message": line.replace("[ERRO]", "Error:").replace("[ERROR]", "Error:")}
        if "[WARN" in line or "[WARNING" in line or "[AVISO]" in line:
            return {"type": "Warning", "message": line}
        return None

    def _read_output(self, process: subprocess.Popen[str]) -> None:
        if process.stdout is None:
            return
        for raw_line in process.stdout:
            line = raw_line.strip()
            if not line:
                continue
            if "[AUTH][PAUSED]" in line or "[RUN_PAUSED]" in line:
                self._paused_for_auth = True
            for marker in ("Run folder:", "Pasta desta execucao:"):
                if marker in line:
                    self.run_dir = line.split(marker, 1)[1].strip()
                    break
            entry = self._actionable_log_entry(line)
            if entry:
                self._logs.append(entry)
        return_code = process.wait()
        if return_code == 0 and self.stop_requested:
            self._logs.append({
                "type": "System",
                "message": "Download pipeline paused safely; pending POs remain queued for resume.",
            })
        else:
            self._logs.append({
                "type": "Success" if return_code == 0 else "Error",
                "message": "Download pipeline finished." if return_code == 0 else f"Download pipeline exited with code {return_code}.",
            })
        self._protect_archived_inputs()

    @staticmethod
    def _set_file_readonly(path: Path, readonly: bool = True) -> None:
        """Toggle the OS read-only attribute on a file (best effort)."""
        try:
            if readonly:
                os.chmod(path, 0o444)
            else:
                os.chmod(path, 0o644)
        except OSError:
            pass

    def _protect_archived_inputs(self) -> None:
        """Make archived inputs read-only after a run finishes so they stay an
        immutable audit snapshot. Explicit retries remove the protection
        temporarily and re-apply it afterwards."""
        session_id = self.session_id
        if not session_id:
            return
        context = self._session_context(session_id)
        input_path = context.get("input_path") or ""
        if input_path and Path(input_path).name.startswith(("input_source_", "input_restored_")):
            self._set_file_readonly(Path(input_path), True)

    def set_run_description(self, session_id: int, description: str) -> dict[str, Any]:
        """Store the free-form run title/description used as audit context."""
        try:
            with self._connect() as conn:
                conn.execute(
                    "UPDATE sessions SET description = ?, updated_at = CURRENT_TIMESTAMP WHERE id = ?",
                    (str(description or "").strip() or None, int(session_id)),
                )
                conn.commit()
            return {"success": True, "session_id": int(session_id)}
        except (sqlite3.Error, ValueError) as exc:
            return {"success": False, "error": str(exc)}

    def _command(
        self,
        *,
        retry_errors: bool = False,
        retry_incomplete: bool = False,
        retry_po: Optional[str] = None,
        retry_in_place_po: Optional[str] = None,
        retry_in_place_errors: bool = False,
        retry_in_place_supplier: Optional[str] = None,
        resume_in_place_session_id: Optional[int] = None,
        provisional_retry_attempt_id: Optional[int] = None,
        retry_staging_dir: Optional[str] = None,
        source_session_id: Optional[int] = None,
        run_dir: Optional[str] = None,
        concurrency: int = 11,
    ) -> list[str]:
        if getattr(sys, "frozen", False):
            command = [sys.executable, "--cli-pipeline", "--concurrency", str(max(1, int(concurrency)))]
        else:
            interpreter = Path(sys.executable)
            # The portable GUI starts with pythonw.exe to avoid a console, but
            # its CLI worker needs python.exe so redirected logs remain usable.
            console_python = interpreter.with_name("python.exe")
            if sys.platform == "win32" and interpreter.name.lower() == "pythonw.exe" and console_python.exists():
                interpreter = console_python
            command = [str(interpreter), str(self.script), "--concurrency", str(max(1, int(concurrency)))]
        if self.download_root:
            command.extend(["--download-root", self.download_root])
        if run_dir:
            command.extend(["--run-dir", run_dir])
        if retry_in_place_po:
            command.extend(["--retry-in-place-po", retry_in_place_po])
            if source_session_id:
                command.extend(["--retry-session-id", str(source_session_id)])
        elif retry_in_place_errors:
            command.append("--retry-in-place-errors")
            if source_session_id:
                command.extend(["--retry-session-id", str(source_session_id)])
        elif retry_in_place_supplier:
            command.extend(["--retry-in-place-supplier", retry_in_place_supplier])
            if source_session_id:
                command.extend(["--retry-session-id", str(source_session_id)])
        elif resume_in_place_session_id:
            command.extend(["--resume-in-place-session-id", str(resume_in_place_session_id)])
        if provisional_retry_attempt_id:
            command.extend([
                "--provisional-retry",
                "--retry-attempt-id",
                str(provisional_retry_attempt_id),
                "--retry-staging-dir",
                str(retry_staging_dir or ""),
            ])
            if source_session_id:
                command.extend(["--retry-session-id", str(source_session_id)])
        if retry_po:
            command.extend(["--retry-po", retry_po])
            if source_session_id:
                command.extend(["--retry-session-id", str(source_session_id)])
        if retry_errors:
            command.append("--retry-last-errors")
            if source_session_id:
                command.extend(["--retry-session-id", str(source_session_id)])
        if retry_incomplete:
            command.extend(["--retry-incomplete-session-id", str(source_session_id)])
        return command

    def start(
        self,
        input_path: str,
        download_root: str,
        concurrency: int = 11,
        *,
        retry_errors: bool = False,
        retry_incomplete: bool = False,
        retry_po: Optional[str] = None,
        retry_in_place_po: Optional[str] = None,
        retry_in_place_errors: bool = False,
        retry_in_place_supplier: Optional[str] = None,
        resume_in_place_session_id: Optional[int] = None,
        provisional_retry_attempt_id: Optional[int] = None,
        retry_staging_dir: Optional[str] = None,
        source_session_id: Optional[int] = None,
        run_dir: Optional[str] = None,
        process_run_dir: Optional[str] = None,
        hierarchy_order: Optional[list[str]] = None,
        retry_attempts: Optional[int] = None,
        msg_processing: str = "convert_extract",
        deduplicate_files: bool = True,
        description: Optional[str] = None,
        column_mapping: Optional[dict[str, str]] = None,
        source_metadata: Optional[dict[str, Any]] = None,
        auth_browser: Optional[str] = None,
    ) -> dict[str, Any]:
        with self._lock:
            if self.process and self.process.poll() is None:
                return {"success": False, "error": "A CLI run is already active."}
            self.input_path = str(Path(input_path).expanduser().resolve())
            self.download_root = str(Path(download_root or self.default_download_root).expanduser().resolve())
            self.source_session_id = source_session_id
            self.run_dir = run_dir
            self.session_id = source_session_id if (retry_in_place_po or retry_in_place_errors or retry_in_place_supplier or resume_in_place_session_id or provisional_retry_attempt_id) else None
            self.stop_requested = False
            self.cancel_requested = False
            self._paused_for_auth = False
            self.started_at = time.time()
            self._logs.clear()
            self._progress_samples.clear()
            self._progress_samples.append((self.started_at, 0))
            self._last_processed = 0
            self._last_progress_at = self.started_at
            try:
                with self._connect() as conn:
                    row = conn.execute("SELECT COALESCE(MAX(id), 0) AS id FROM sessions").fetchone()
                    self.previous_max_session_id = int(row["id"] or 0)
            except sqlite3.Error:
                self.previous_max_session_id = 0

            env = os.environ.copy()
            env["INPUT_CSV"] = self.input_path
            if hierarchy_order is not None:
                env["COUPA_HIERARCHY_ORDER"] = json.dumps(hierarchy_order)
            if description:
                env["COUPA_RUN_DESCRIPTION"] = str(description).strip()
            if source_metadata:
                env["COUPA_SOURCE_METADATA"] = json.dumps(source_metadata, ensure_ascii=False)
            if column_mapping:
                env["COUPA_COLUMN_MAPPING"] = json.dumps(column_mapping)
            env["COUPA_RETRY_ATTEMPTS"] = str(max(1, min(3, int(retry_attempts or 1))))
            env["COUPA_MSG_PROCESSING"] = msg_processing if msg_processing in {"disabled", "convert", "convert_extract"} else "convert_extract"
            env["COUPA_DEDUPLICATE_FILES"] = "1" if deduplicate_files else "0"
            # GUI authentication is completed before the worker starts. The
            # worker validates the shared cache but never opens an untracked
            # second browser; direct CLI invocation keeps its own interactive
            # default because it does not pass this environment flag.
            env["COUPA_AUTH_INTERACTIVE"] = "0"
            if auth_browser in {"auto", "edge", "chrome"}:
                env["COUPA_AUTH_BROWSER"] = auth_browser
            command = self._command(
                retry_errors=retry_errors,
                retry_incomplete=retry_incomplete,
                retry_po=retry_po,
                retry_in_place_po=retry_in_place_po,
                retry_in_place_errors=retry_in_place_errors,
                retry_in_place_supplier=retry_in_place_supplier,
                resume_in_place_session_id=source_session_id if resume_in_place_session_id else None,
                provisional_retry_attempt_id=provisional_retry_attempt_id,
                retry_staging_dir=retry_staging_dir,
                source_session_id=source_session_id,
                run_dir=process_run_dir or run_dir,
                concurrency=concurrency,
            )
            try:
                self.process = subprocess.Popen(
                    command,
                    cwd=str(self.project_root),
                    env=env,
                    stdin=subprocess.DEVNULL,
                    stdout=subprocess.PIPE,
                    stderr=subprocess.STDOUT,
                    text=True,
                    bufsize=1,
                    creationflags=self._process_creationflags(),
                )
            except OSError as exc:
                self.process = None
                return {"success": False, "error": str(exc)}
            threading.Thread(target=self._read_output, args=(self.process,), daemon=True).start()
            return {"success": True, "session_id": 0, "message": "CLI pipeline started."}

    def _resolve_session_id(self) -> Optional[int]:
        if self.session_id:
            return self.session_id
        if self.source_session_id and self.process and self.process.poll() is None:
            latest = self._latest_session_id()
            if latest:
                self.session_id = latest
        else:
            latest = self._latest_session_id()
            if latest:
                self.session_id = latest
        if self.session_id:
            self._remember_session(self.session_id)
        return self.session_id

    def _stats(self, session_id: int) -> dict[str, int]:
        with self._connect() as conn:
            row = conn.execute(
                """
                SELECT COUNT(*) AS total,
                       SUM(CASE WHEN status = 'SUCCESS' THEN 1 ELSE 0 END) AS success,
                       SUM(CASE WHEN status IN ('ERROR', 'SKIPPED_VERIFICATION_REQUIRED') THEN 1 ELSE 0 END) AS errors,
                       SUM(CASE WHEN status = 'PENDING' THEN 1 ELSE 0 END) AS pending
                FROM po_downloads WHERE session_id = ?
                """,
                (session_id,),
            ).fetchone()
        return {key: int(row[key] or 0) for key in ("total", "success", "errors", "pending")}

    def _status(self, stats: dict[str, int]) -> str:
        running = bool(self.process and self.process.poll() is None)
        if running:
            return "STOPPING" if self.stop_requested else "RUNNING"
        if self.cancel_requested:
            return "CANCELLED"
        if self._paused_for_auth:
            return "RUN_PAUSED"
        if self.stop_requested:
            return "STOPPED"
        if stats["pending"] > 0:
            return "PARTIAL" if stats["success"] else "FAILED"
        if stats["errors"] == 0:
            return "SUCCESS"
        return "PARTIAL" if stats["success"] else "FAILED"

    @staticmethod
    def _format_eta(seconds: int) -> str:
        seconds = max(0, int(seconds))
        if seconds >= 3600:
            hours, remainder = divmod(seconds, 3600)
            minutes = remainder // 60
            return f"{hours}h {minutes:02d}m"
        minutes, remaining_seconds = divmod(seconds, 60)
        return f"{minutes:02d}:{remaining_seconds:02d}"

    def _recent_progress_metrics(self, processed: int, total: int, now: Optional[float] = None) -> tuple[float, str, int]:
        """Return rolling 60-second throughput, ETA, and seconds since progress."""
        current_time = time.time() if now is None else now
        if processed < self._last_processed:
            self._progress_samples.clear()
            self._progress_samples.append((current_time, processed))
        if processed > self._last_processed:
            self._last_progress_at = current_time
        self._last_processed = processed
        self._progress_samples.append((current_time, processed))

        cutoff = current_time - 60.0
        samples = list(self._progress_samples)
        candidates = [sample for sample in samples if sample[0] >= cutoff]
        if samples and candidates and candidates[0] != samples[0]:
            previous = [sample for sample in samples if sample[0] < cutoff]
            if previous and cutoff - previous[-1][0] <= 5.0:
                candidates.insert(0, previous[-1])
        if not candidates:
            candidates = samples[-1:]

        speed = 0.0
        if len(candidates) >= 2:
            first_time, first_processed = candidates[0]
            elapsed = current_time - first_time
            completed = processed - first_processed
            if elapsed >= 5.0 and completed > 0:
                speed = completed / elapsed * 60.0

        stalled_seconds = max(0, int(current_time - self._last_progress_at)) if processed < total else 0
        if stalled_seconds >= 45:
            speed = 0.0
        remaining = max(total - processed, 0)
        if not remaining:
            eta = "00:00"
        elif speed > 0:
            eta = self._format_eta(round(remaining / speed * 60.0))
        else:
            eta = "--:--"
        return speed, eta, stalled_seconds

    def get_status(self, requested_session_id: int = 0) -> dict[str, Any]:
        with self._lock:
            session_id = self._resolve_session_id()
            if not session_id:
                process_finished = self.process is not None and self.process.poll() is not None
                return {
                    "status": "FAILED" if process_finished else "STARTING",
                    "session_id": 0,
                    "total": 0,
                    "processed": 0,
                    "success": 0,
                    "errors": 0,
                    "speed": 0.0,
                    "eta": "--:--",
                    "latest_logs": list(self._drain_logs()),
                }
            try:
                stats = self._stats(session_id)
            except sqlite3.Error as exc:
                return {"status": "FAILED", "session_id": session_id, "total": 0, "processed": 0, "success": 0, "errors": 1, "speed": 0.0, "eta": "--:--", "latest_logs": [{"type": "Error", "message": str(exc)}]}
            processed = stats["success"] + stats["errors"]
            speed, eta, stalled_seconds = self._recent_progress_metrics(processed, stats["total"])
            return {
                "status": self._status(stats),
                "session_id": session_id,
                "total": stats["total"],
                "processed": processed,
                "success": stats["success"],
                "errors": stats["errors"],
                "speed": speed,
                "eta": eta,
                "speed_window_seconds": 60,
                "stalled_seconds": stalled_seconds,
                "run_dir": self.run_dir or "",
                "latest_logs": list(self._drain_logs()),
            }

    def _drain_logs(self) -> list[dict[str, str]]:
        logs = list(self._logs)
        self._logs.clear()
        return logs

    def _interrupt(self, *, resumable: bool) -> dict[str, Any]:
        if not self.process or self.process.poll() is not None:
            return {"success": False, "error": "No active CLI run."}
        self.stop_requested = resumable
        self.cancel_requested = not resumable
        try:
            if os.name == "nt":
                self.process.send_signal(signal.CTRL_BREAK_EVENT)
            else:
                self.process.send_signal(signal.SIGINT)
            return {"success": True}
        except OSError as exc:
            return {"success": False, "error": str(exc)}

    def pause(self) -> dict[str, Any]:
        return self._interrupt(resumable=True)

    def stop(self) -> dict[str, Any]:
        return self._interrupt(resumable=False)

    def reconcile(self, session_id: int) -> dict[str, Any]:
        """Mark successful rows with missing files as PENDING before resume."""
        repaired = []
        with self._connect() as conn:
            rows = conn.execute(
                "SELECT po_number, status, attachment_count, download_folder, output_subdir FROM po_downloads WHERE session_id = ?",
                (session_id,),
            ).fetchall()
            for row in rows:
                folder_value = str(row["download_folder"] or "").strip()
                if not folder_value and self.run_dir and row["output_subdir"]:
                    folder_value = str(Path(self.run_dir) / str(row["output_subdir"]) / str(row["po_number"]))
                folder = Path(folder_value) if folder_value else Path("__missing_folder__")
                expected = int(row["attachment_count"] or 0)
                actual = sum(1 for path in folder.iterdir() if path.is_file() and path.name != ".DS_Store") if folder.exists() else 0
                skipped = row["status"] == "SKIPPED_VERIFICATION_REQUIRED"
                inconsistent_success = row["status"] == "SUCCESS" and expected > 0 and actual < expected
                unfinished_with_files = row["status"] in {"PENDING", "ERROR"} and actual > 0
                if skipped:
                    conn.execute(
                        "UPDATE po_downloads SET status = 'PENDING', attachment_count = 0, error_message = ? WHERE session_id = ? AND po_number = ?",
                        ("Supplier verification was previously required; scheduled for retry.", session_id, row["po_number"]),
                    )
                    repaired.append(row["po_number"])
                    continue
                if not inconsistent_success and not unfinished_with_files:
                    continue
                if folder.exists():
                    backup = folder.with_name(f"{folder.name}.partial_backup_{time.strftime('%Y%m%d-%H%M%S')}")
                    shutil.move(str(folder), str(backup))
                message = f"Reconciliation found {actual}/{expected} files; scheduled for clean retry."
                conn.execute(
                    "UPDATE po_downloads SET status = 'PENDING', attachment_count = 0, download_folder = NULL, error_message = ? WHERE session_id = ? AND po_number = ?",
                    (message, session_id, row["po_number"]),
                )
                repaired.append(row["po_number"])
            conn.commit()
        return {"success": True, "repaired": repaired, "count": len(repaired)}

    def resume(self) -> dict[str, Any]:
        source = self.session_id
        if not source or not self.input_path:
            return {"success": False, "error": "No paused session found."}
        reconciliation = self.reconcile(source)
        return {
            **self.start(
                self.input_path,
                self.download_root or str(self.default_download_root),
                resume_in_place_session_id=source,
                source_session_id=source,
                run_dir=self.run_dir,
            ),
            "reconciliation": reconciliation,
        }

    @staticmethod
    def _run_dir_from_download_folder(download_folder: str, output_subdir: str = "") -> str:
        folder = Path(download_folder)
        subdir_parts = [part for part in Path(output_subdir or "").parts if part not in {"", "."}]
        candidate = folder.parent
        for _ in range(len(subdir_parts)):
            candidate = candidate.parent
        return str(candidate)

    def _session_context(self, session_id: int, po_number: str | None = None) -> dict[str, str]:
        metadata = self._session_metadata.get(str(session_id), {}).copy()
        with self._connect() as conn:
            session = conn.execute("SELECT * FROM sessions WHERE id = ?", (session_id,)).fetchone()
            if po_number:
                row = conn.execute(
                    "SELECT download_folder, output_subdir FROM po_downloads WHERE session_id = ? AND po_number = ?",
                    (session_id, po_number),
                ).fetchone()
                reference_row = row
                if not reference_row or not reference_row["download_folder"]:
                    reference_row = conn.execute(
                        """
                        SELECT download_folder, output_subdir
                        FROM po_downloads
                        WHERE session_id = ? AND download_folder IS NOT NULL AND download_folder != ''
                        ORDER BY id ASC LIMIT 1
                        """,
                        (session_id,),
                    ).fetchone()
            else:
                row = None
                reference_row = None
        if session:
            stored_path = str(session["input_file_path"] or "")
            if stored_path and Path(stored_path).exists():
                metadata["input_path"] = stored_path
            elif not metadata.get("input_path"):
                metadata["input_path"] = stored_path or str(session["input_file"] or "")
            if not metadata.get("run_dir") and stored_path:
                stored_name = Path(stored_path).name
                if stored_name.startswith(("input_source_", "input_restored_")):
                    metadata["run_dir"] = str(Path(stored_path).parent)
        if reference_row and not metadata.get("run_dir") and reference_row["download_folder"]:
            metadata["run_dir"] = self._run_dir_from_download_folder(
                str(reference_row["download_folder"]),
                str(reference_row["output_subdir"] or ""),
            )
        if not metadata.get("run_dir"):
            roots = []
            if self.download_root:
                roots.append(Path(self.download_root))
            roots.append(self.default_download_root)
            report_name = f"report_session_{session_id}.xlsx"
            for root in roots:
                matches = list(root.glob(f"run_*/{report_name}")) if root.exists() else []
                if matches:
                    metadata["run_dir"] = str(matches[0].parent)
                    break
        if session and session["input_file_blob"] and metadata.get("run_dir"):
            current = Path(str(metadata.get("input_path") or ""))
            if not current.exists():
                suffix = Path(str(session["input_file"] or "")).suffix or ".csv"
                restored = Path(metadata["run_dir"]) / f"input_restored_{session_id}{suffix}"
                try:
                    restored.parent.mkdir(parents=True, exist_ok=True)
                    restored.write_bytes(session["input_file_blob"])
                    with self._connect() as conn:
                        conn.execute(
                            "UPDATE sessions SET input_file_path = ?, updated_at = CURRENT_TIMESTAMP WHERE id = ?",
                            (str(restored), session_id),
                        )
                        conn.commit()
                    metadata["input_path"] = str(restored)
                except OSError:
                    pass
        return metadata

    @staticmethod
    def _detail_header(value: Any) -> str:
        return re.sub(r"[^a-z0-9]+", "_", str(value or "").strip().casefold()).strip("_")

    @classmethod
    def _supplier_by_po_from_input(cls, input_path: str) -> dict[str, str]:
        """Read supplier labels for run summaries without changing the run schema."""
        path = Path(str(input_path or "")).expanduser()
        if not path.is_file():
            return {}
        try:
            if path.suffix.casefold() in {".csv", ".tsv", ".txt"}:
                with path.open("r", encoding="utf-8-sig", newline="") as handle:
                    sample = handle.read(4096)
                    handle.seek(0)
                    try:
                        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
                    except csv.Error:
                        dialect = csv.excel_tab if "\t" in sample else csv.excel
                    rows = csv.DictReader(handle, dialect=dialect)
                    records = list(rows)
            else:
                from openpyxl import load_workbook

                sheet = load_workbook(path, read_only=True, data_only=True).active
                values = sheet.iter_rows(values_only=True)
                headers = next(values, ())
                records = [dict(zip(headers, row)) for row in values]
            if not records:
                return {}
            headers = list(records[0].keys())
            normalized = {header: cls._detail_header(header) for header in headers}
            po_header = next((header for header, key in normalized.items() if key in {"po", "po_number", "ponumber", "purchase_order"}), None)
            supplier_header = next((header for header, key in normalized.items() if key in {"supplier", "supplier_uu", "supplier_uu_name", "supplier_uu_code"} or "supplier" in key), None)
            if not po_header or not supplier_header:
                return {}
            result = {}
            for record in records:
                po = str(record.get(po_header) or "").strip()
                supplier = str(record.get(supplier_header) or "").strip()
                if po and supplier:
                    result[po] = supplier
            return result
        except (OSError, ValueError, KeyError, csv.Error, BadZipFile):
            return {}

    def open_input_file(self, session_id: int) -> dict[str, Any]:
        context = self._session_context(session_id)
        path = Path(context.get("input_path", "")).expanduser()
        if not path.exists():
            return {"success": False, "error": "The preserved input file is not available for this run."}
        return self._open_path(path)

    @staticmethod
    def _open_path(path: Path) -> dict[str, Any]:
        """Open a local file or folder with the operating system launcher."""
        try:
            if sys.platform == "darwin":
                subprocess.Popen(["open", str(path)])
            elif os.name == "nt":
                os.startfile(str(path))  # type: ignore[attr-defined]
            else:
                subprocess.Popen(["xdg-open", str(path)])
            return {"success": True, "path": str(path)}
        except OSError as exc:
            return {"success": False, "error": str(exc)}

    def open_run_folder(self, session_id: int) -> dict[str, Any]:
        context = self._session_context(int(session_id))
        run_dir_value = str(context.get("run_dir") or "").strip()
        run_dir = Path(run_dir_value).expanduser() if run_dir_value else None
        if run_dir is None or not run_dir.is_dir():
            return {"success": False, "error": "The download folder is not available for this run."}
        return self._open_path(run_dir)

    def open_run_report(self, session_id: int) -> dict[str, Any]:
        exported = self.export_report(int(session_id))
        if not exported.get("success"):
            return exported
        report_path = Path(str(exported.get("filepath", ""))).expanduser()
        if not report_path.is_file():
            return {"success": False, "error": "The Excel report is not available for this run."}
        return self._open_path(report_path)

    def retry_po_with_edit(self, session_id: int, original_po: str, edited_po: str) -> dict[str, Any]:
        """Run an edited PO in a staging directory until the user commits it."""
        original = str(original_po or "").strip()
        edited = str(edited_po or "").strip()
        if not original or not edited:
            return {"success": False, "error": "A PO number is required."}
        if original == edited:
            return self.retry_po(session_id, original)
        if not re.fullmatch(r"(?i)(?:PO|PM)[0-9]+", edited):
            return {"success": False, "error": "Enter a valid PO number using PO or PM followed by digits."}

        context = self._session_context(session_id, original)
        run_dir = context.get("run_dir") or self.run_dir
        input_path = context.get("input_path") or self.input_path
        if not run_dir or not input_path or not Path(input_path).exists():
            return {"success": False, "error": "The original run files are not available."}

        attempt_token = uuid.uuid4().hex[:12]
        staging_dir = Path(run_dir) / ".retry_staging" / f"attempt_{attempt_token}"
        attempt_id = 0
        try:
            with self._connect() as conn:
                source = conn.execute(
                    "SELECT po_number, company_code, output_subdir, status FROM po_downloads WHERE session_id = ? AND po_number = ?",
                    (session_id, original),
                ).fetchone()
                if not source:
                    return {"success": False, "error": "The original PO was not found in this run."}
                duplicate = conn.execute(
                    "SELECT 1 FROM po_downloads WHERE session_id = ? AND po_number = ?",
                    (session_id, edited),
                ).fetchone()
                if duplicate:
                    return {"success": False, "error": "The corrected PO already exists in this run."}
                staging_dir.mkdir(parents=True, exist_ok=True)
                cursor = conn.execute(
                    "INSERT INTO retry_attempts (session_id, original_po_number, edited_po_number, staging_dir) VALUES (?, ?, ?, ?)",
                    (session_id, original, edited, str(staging_dir)),
                )
                attempt_id = int(cursor.lastrowid)
                conn.execute(
                    "INSERT INTO po_downloads (session_id, po_number, company_code, output_subdir, status) VALUES (?, ?, ?, ?, 'PENDING')",
                    (session_id, edited, source["company_code"], source["output_subdir"]),
                )
                conn.commit()

            self.input_path = str(Path(input_path).expanduser().resolve())
            self.run_dir = str(Path(run_dir).expanduser().resolve())
            result = self.start(
                self.input_path,
                str(Path(run_dir).parent),
                retry_attempts=1,
                source_session_id=session_id,
                run_dir=self.run_dir,
                process_run_dir=str(staging_dir),
                provisional_retry_attempt_id=attempt_id,
                retry_staging_dir=str(staging_dir),
                msg_processing="disabled",
                deduplicate_files=False,
            )
            if not result.get("success"):
                self.discard_retry_attempt(attempt_id)
                return result
            return {**result, "session_id": session_id, "attempt_id": attempt_id, "original_po": original, "edited_po": edited}
        except (OSError, sqlite3.Error) as exc:
            if attempt_id:
                self.discard_retry_attempt(attempt_id)
            return {"success": False, "error": str(exc)}

    @staticmethod
    def _po_values_equal(left: Any, right: Any) -> bool:
        def normalize(value: Any) -> str:
            text = str(value or "").strip().upper()
            text = re.sub(r"\.0+$", "", text)
            if text.startswith(("PO", "PM")):
                text = text[2:]
            return text

        return normalize(left) == normalize(right)

    @staticmethod
    def _replace_po_in_file(path: Path, old_po: str, new_po: str, note: str) -> None:
        if not path.exists():
            raise FileNotFoundError(str(path))
        if path.suffix.lower() in {".xlsx", ".xlsm"}:
            from openpyxl import load_workbook
            workbook = load_workbook(path)
            sheet = workbook["Input"] if "Input" in workbook.sheetnames else workbook.active
            headers = [str(cell.value or "").strip() for cell in sheet[1]]
            po_index = next((index for index, value in enumerate(headers, 1) if re.sub(r"[^a-z0-9]+", "", value.lower()) in {"ponumber", "po", "pedido"}), None)
            if not po_index:
                raise ValueError(f"PO_NUMBER column not found in {path.name}")
            remarks_index = next((index for index, value in enumerate(headers, 1) if re.sub(r"[^a-z0-9]+", "", value.lower()) in {"remarks", "observations", "observacao", "observacoes"}), None)
            if not remarks_index:
                remarks_index = sheet.max_column + 1
                sheet.cell(1, remarks_index).value = "REMARKS"
            found = False
            already_updated = False
            for row in range(2, sheet.max_row + 1):
                value = sheet.cell(row, po_index).value
                if CliProcessSupervisor._po_values_equal(value, old_po):
                    sheet.cell(row, po_index).value = new_po
                    sheet.cell(row, remarks_index).value = note
                    found = True
                elif CliProcessSupervisor._po_values_equal(value, new_po):
                    sheet.cell(row, remarks_index).value = note
                    already_updated = True
            if not found and not already_updated:
                raise ValueError(f"PO {old_po} was not found in {path.name}")
            workbook.save(path)
            return
        import pandas as pd
        with open(path, encoding="utf-8-sig", newline="") as handle:
            sample = handle.read(4096)
        separator = ";" if sample.count(";") > sample.count(",") else ","
        frame = pd.read_csv(path, sep=separator, dtype=str, encoding="utf-8-sig").fillna("")
        po_column = next((column for column in frame.columns if re.sub(r"[^a-z0-9]+", "", str(column).lower()) in {"ponumber", "po", "pedido"}), None)
        if po_column is None:
            raise ValueError(f"PO_NUMBER column not found in {path.name}")
        values = frame[po_column].astype(str).str.strip()
        mask = values.map(lambda value: CliProcessSupervisor._po_values_equal(value, old_po))
        already_updated = values.map(lambda value: CliProcessSupervisor._po_values_equal(value, new_po))
        if not mask.any() and not already_updated.any():
            raise ValueError(f"PO {old_po} was not found in {path.name}")
        frame.loc[mask, po_column] = new_po
        mask = mask | already_updated
        remarks_column = next((column for column in frame.columns if re.sub(r"[^a-z0-9]+", "", str(column).lower()) in {"remarks", "observations", "observacao", "observacoes"}), "REMARKS")
        if remarks_column not in frame.columns:
            frame[remarks_column] = ""
        frame.loc[mask, remarks_column] = note
        frame.to_csv(path, sep=separator, index=False, encoding="utf-8-sig")

    def persist_retry_files(self, attempt_id: int) -> dict[str, Any]:
        """Persist a committed retry in the archived input, original input, and report."""
        with self._connect() as conn:
            attempt = conn.execute("SELECT * FROM retry_attempts WHERE id = ? AND status = 'COMMITTED'", (int(attempt_id),)).fetchone()
            session = conn.execute("SELECT input_file_path FROM sessions WHERE id = ?", (attempt["session_id"],)).fetchone() if attempt else None
        if not attempt or not session:
            return {"success": False, "error": "Committed retry metadata was not found."}
        note = str(attempt["error_message"] or f"Corrected from {attempt['original_po_number']} to {attempt['edited_po_number']}; retry succeeded.")
        archive = Path(str(session["input_file_path"] or ""))
        metadata = self._session_metadata.get(str(attempt["session_id"]), {})
        original_value = metadata.get("source_input_path") or metadata.get("input_path") or ""
        original = Path(str(original_value)).expanduser() if original_value else None
        report_dir = Path(self._session_context(int(attempt["session_id"])).get("run_dir") or self.run_dir or "")
        report = report_dir / f"report_session_{attempt['session_id']}.xlsx"
        required_paths = [archive, original, report]
        missing = [str(path) for path in required_paths if not path or not path.exists()]
        if missing:
            return {"success": False, "error": f"Could not find files to update: {', '.join(missing)}"}
        paths: list[Path] = []
        for path in required_paths:
            if path not in paths:
                paths.append(path)
        try:
            backups: list[tuple[Path, Path]] = []
            for path in paths:
                backup = path.with_name(f"{path.stem}.before_retry_{time.strftime('%Y%m%d-%H%M%S')}{path.suffix}")
                shutil.copy2(path, backup)
                backups.append((path, backup))
            # The archived input is read-only between runs; an explicit retry
            # is the only flow allowed to update it, so lift the protection
            # for the duration of the edit and re-apply it afterwards.
            for path in paths:
                self._set_file_readonly(path, False)
            for path in paths:
                self._replace_po_in_file(path, str(attempt["original_po_number"]), str(attempt["edited_po_number"]), note)
            if archive.exists():
                data = archive.read_bytes()
                with self._connect() as conn:
                    conn.execute(
                        "UPDATE sessions SET input_file_path = ?, input_file_blob = ?, input_file_sha256 = ?, input_file_size = ?, updated_at = CURRENT_TIMESTAMP WHERE id = ?",
                        (str(archive), data, hashlib.sha256(data).hexdigest(), len(data), int(attempt["session_id"])),
                    )
                    conn.commit()
                self._set_file_readonly(archive, True)
            return {"success": True, "paths": [str(path) for path in paths], "note": note}
        except (OSError, ValueError, sqlite3.Error) as exc:
            for path, backup in locals().get("backups", []):
                try:
                    shutil.copy2(backup, path)
                except OSError:
                    pass
            return {"success": False, "error": f"Could not persist retry correction: {exc}"}

    def get_retry_attempt_status(self, attempt_id: int) -> dict[str, Any]:
        with self._connect() as conn:
            row = conn.execute("SELECT * FROM retry_attempts WHERE id = ?", (int(attempt_id),)).fetchone()
            if not row:
                return {"success": False, "error": "Retry attempt not found."}
            target = conn.execute(
                "SELECT status, attachment_count, error_message, download_folder FROM po_downloads WHERE session_id = ? AND po_number = ?",
                (row["session_id"], row["edited_po_number"]),
            ).fetchone()
        status = str(row["status"] or "RUNNING")
        target_status = str(target["status"] if target else "PENDING")
        process_running = bool(self.process and self.process.poll() is None)
        if process_running:
            status = "RUNNING"
        elif status == "RUNNING":
            status = "SUCCESS" if target_status == "SUCCESS" else "FAILED"
        return {
            "success": True,
            "attempt_id": int(row["id"]),
            "session_id": int(row["session_id"]),
            "original_po": row["original_po_number"],
            "edited_po": row["edited_po_number"],
            "status": status,
            "target_status": target_status,
            "attachment_count": int((target["attachment_count"] if target else 0) or 0),
            "error_message": (target["error_message"] if target else None) or row["error_message"],
            "download_folder": target["download_folder"] if target else None,
            "staging_dir": row["staging_dir"],
        }

    def discard_retry_attempt(self, attempt_id: int) -> dict[str, Any]:
        with self._lock:
            if self.process and self.process.poll() is None:
                return {"success": False, "error": "Wait for the retry to finish before discarding it."}
        try:
            with self._connect() as conn:
                row = conn.execute("SELECT * FROM retry_attempts WHERE id = ?", (int(attempt_id),)).fetchone()
                if not row:
                    return {"success": False, "error": "Retry attempt not found."}
                conn.execute(
                    "DELETE FROM po_downloads WHERE session_id = ? AND po_number = ?",
                    (row["session_id"], row["edited_po_number"]),
                )
                conn.execute(
                    "UPDATE retry_attempts SET status = 'DISCARDED', completed_at = CURRENT_TIMESTAMP WHERE id = ?",
                    (int(attempt_id),),
                )
                conn.commit()
            shutil.rmtree(str(row["staging_dir"]), ignore_errors=True)
            return {"success": True, "session_id": int(row["session_id"])}
        except (OSError, sqlite3.Error) as exc:
            return {"success": False, "error": str(exc)}

    def commit_retry_attempt(self, attempt_id: int) -> dict[str, Any]:
        with self._lock:
            if self.process and self.process.poll() is None:
                return {"success": False, "error": "Wait for the retry to finish before saving it."}
        try:
            with self._connect() as conn:
                attempt = conn.execute("SELECT * FROM retry_attempts WHERE id = ?", (int(attempt_id),)).fetchone()
                if not attempt:
                    return {"success": False, "error": "Retry attempt not found."}
                target = conn.execute(
                    "SELECT * FROM po_downloads WHERE session_id = ? AND po_number = ?",
                    (attempt["session_id"], attempt["edited_po_number"]),
                ).fetchone()
                if not target or target["status"] != "SUCCESS":
                    return {"success": False, "error": "Only a successful retry can be saved."}
                source = conn.execute(
                    "SELECT * FROM po_downloads WHERE session_id = ? AND po_number = ?",
                    (attempt["session_id"], attempt["original_po_number"]),
                ).fetchone()
                if not source:
                    return {"success": False, "error": "The original PO record was not found."}
                final_dir = Path(self._session_context(int(attempt["session_id"])).get("run_dir") or self.run_dir or "")
                staged = Path(str(attempt["staging_dir"]))
                source_folder = Path(str(source["download_folder"] or ""))
                target_folder = final_dir / str(target["output_subdir"] or source["output_subdir"] or target["company_code"]) / str(target["po_number"])
                staged_folder = staged / str(target["output_subdir"] or target["company_code"]) / str(target["po_number"])
                if staged_folder.exists():
                    target_folder.parent.mkdir(parents=True, exist_ok=True)
                    if target_folder.exists():
                        shutil.rmtree(target_folder)
                    shutil.move(str(staged_folder), str(target_folder))
                if str(source["download_folder"] or "").strip() and source_folder.exists() and source_folder != target_folder:
                    shutil.rmtree(source_folder)
                note = f"Corrected from {attempt['original_po_number']} to {attempt['edited_po_number']}; retry succeeded."
                conn.execute("DELETE FROM po_downloads WHERE id = ?", (target["id"],))
                conn.execute(
                    "UPDATE po_downloads SET po_number = ?, status = 'SUCCESS', download_folder = ?, attachment_count = ?, error_message = NULL, remarks = ?, updated_at = CURRENT_TIMESTAMP WHERE id = ?",
                    (attempt["edited_po_number"], str(target_folder), target["attachment_count"], note, source["id"]),
                )
                conn.execute(
                    "UPDATE retry_attempts SET status = 'COMMITTED', completed_at = CURRENT_TIMESTAMP, error_message = ? WHERE id = ?",
                    (note, int(attempt_id)),
                )
                conn.execute(
                    "INSERT INTO retry_events (session_id, po_number, status_before, status_after, error_message, completed_at) VALUES (?, ?, ?, ?, ?, CURRENT_TIMESTAMP)",
                    (attempt["session_id"], attempt["edited_po_number"], source["status"], "SUCCESS", note),
                )
                conn.commit()
            shutil.rmtree(str(attempt["staging_dir"]), ignore_errors=True)
            return {"success": True, "session_id": int(attempt["session_id"]), "po_number": attempt["edited_po_number"], "note": note}
        except (OSError, sqlite3.Error) as exc:
            return {"success": False, "error": str(exc)}

    def retry_po(self, session_id: int, po_number: str) -> dict[str, Any]:
        normalized_po = str(po_number or "").strip()
        if not normalized_po:
            return {"success": False, "error": "PO number is required."}
        context = self._session_context(session_id, normalized_po)
        run_dir = context.get("run_dir") or self.run_dir
        if not run_dir:
            return {"success": False, "error": "The original run folder could not be located for this PO."}
        input_path = context.get("input_path") or self.input_path or "input.csv"
        self.input_path = input_path
        self.run_dir = run_dir
        return self.start(
            input_path,
            self.download_root or str(Path(run_dir).parent),
            retry_in_place_po=normalized_po,
            source_session_id=session_id,
            run_dir=run_dir,
        )

    def retry_errors(self, session_id: int) -> dict[str, Any]:
        context = self._session_context(session_id)
        if not self.input_path:
            self.input_path = context.get("input_path") or None
            self.run_dir = self.run_dir or context.get("run_dir") or None
        if not self.input_path or not Path(self.input_path).exists():
            return {"success": False, "error": "The original input file is not available for this run."}
        return self.start(
            self.input_path,
            self.download_root or str(self.default_download_root),
            retry_in_place_errors=True,
            source_session_id=session_id,
            run_dir=self.run_dir,
        )

    def retry_supplier(self, session_id: int, supplier: str) -> dict[str, Any]:
        value = str(supplier or "").strip()
        if not value:
            return {"success": False, "error": "Supplier is required."}
        context = self._session_context(session_id)
        input_path = context.get("input_path") or self.input_path
        run_dir = context.get("run_dir") or self.run_dir
        if not input_path or not Path(input_path).exists():
            return {"success": False, "error": "The original input file is not available for this run."}
        return self.start(
            input_path,
            self.download_root or str(self.default_download_root),
            retry_in_place_supplier=value,
            source_session_id=session_id,
            run_dir=run_dir,
        )

    @staticmethod
    def _coupa_url(po_number: str) -> str:
        value = str(po_number or "").strip()
        order_number = value[2:] if value.upper().startswith(("PO", "PM")) else value
        return f"https://unilever.coupahost.com/order_headers/{order_number}"

    @staticmethod
    def _derived_status(stats: dict[str, int], stored: str = "PENDING") -> str:
        if stats["pending"] == 0:
            return "SUCCESS" if stats["errors"] == 0 else ("PARTIAL" if stats["success"] else "FAILED")
        return stored or "PENDING"

    def _apply_retention(self) -> None:
        settings_path = Path.home() / ".contract_downloader" / "gui_settings.json"
        try:
            settings = json.loads(settings_path.read_text(encoding="utf-8"))
        except (OSError, ValueError):
            return
        retention = str(settings.get("retention", "all"))
        if retention == "all":
            return
        with self._connect() as conn:
            if retention in {"10", "30"}:
                rows = conn.execute(
                    "SELECT id FROM sessions WHERE status NOT IN ('RUNNING', 'PENDING') ORDER BY id DESC"
                ).fetchall()
                ids = [int(row["id"]) for row in rows[int(retention):]]
            elif retention == "90":
                rows = conn.execute(
                    "SELECT id FROM sessions WHERE status NOT IN ('RUNNING', 'PENDING') AND created_at < datetime('now', '-90 days')"
                ).fetchall()
                ids = [int(row["id"]) for row in rows]
            else:
                return
        for session_id in ids:
            self.delete_session(session_id)

    def concurrency_estimates(self) -> dict[str, dict[str, Any]]:
        """Estimate 100-PO duration from completed local runs."""
        estimates = {str(value): {"minutes_100": None, "samples": 0} for value in (2, 4, 6, 8)}
        with self._connect() as conn:
            rows = conn.execute(
                """
                SELECT concurrency, duration_seconds, COUNT(*) AS total_pos
                FROM sessions s
                LEFT JOIN po_downloads p ON p.session_id = s.id
                WHERE s.concurrency IN (2, 4, 6, 8)
                  AND s.duration_seconds > 0
                GROUP BY s.id
                ORDER BY s.id DESC
                """
            ).fetchall()
        by_concurrency: dict[int, list[float]] = {2: [], 4: [], 6: [], 8: []}
        for row in rows:
            total_pos = int(row["total_pos"] or 0)
            duration = float(row["duration_seconds"] or 0)
            if total_pos > 0 and duration > 0:
                by_concurrency[int(row["concurrency"])].append(100.0 * duration / 60.0 / total_pos)
        for concurrency, values in by_concurrency.items():
            if values:
                values = sorted(values[:10])
                median = values[len(values) // 2]
                estimates[str(concurrency)] = {
                    "minutes_100": round(median, 1),
                    "samples": len(values),
                }
        return estimates

    def history(self) -> list[dict[str, Any]]:
        self._apply_retention()
        with self._connect() as conn:
            sessions = conn.execute("SELECT * FROM sessions ORDER BY id DESC").fetchall()
        result = []
        for session in sessions:
            session_id = int(session["id"])
            stats = self._stats(session_id)
            context = self._session_context(session_id)
            session_data = dict(session)
            session_data.pop("input_file_blob", None)
            run_dir = str(context.get("run_dir") or "")
            report_path = str(Path(run_dir) / f"report_session_{session_id}.xlsx") if run_dir else ""
            result.append({
                **session_data,
                "input_file_path": session["input_file_path"] or context.get("input_path", ""),
                "run_dir": run_dir,
                "report_path": report_path if Path(report_path).is_file() else "",
                "status": self._derived_status(stats, session["status"]),
                "total_pos": stats["total"],
                "success_count": stats["success"],
                "error_count": stats["errors"],
                "pending_count": stats["pending"],
            })
        return result

    def reset_local_state_preserving_files(self) -> dict[str, Any]:
        """Forget local run records without deleting downloaded files."""
        with self._lock:
            if self.process and self.process.poll() is None:
                return {"success": False, "error": "Stop the active run before starting clean."}
        try:
            with self._connect() as conn:
                conn.execute("DELETE FROM retry_events")
                conn.execute("DELETE FROM retry_attempts")
                conn.execute("DELETE FROM po_attachments")
                conn.execute("DELETE FROM po_downloads")
                conn.execute("DELETE FROM sessions")
                try:
                    conn.execute("DELETE FROM sqlite_sequence WHERE name IN ('sessions', 'po_downloads', 'retry_events', 'retry_attempts')")
                except sqlite3.OperationalError:
                    pass
                conn.commit()
        except sqlite3.Error as exc:
            return {"success": False, "error": f"Could not reset local run state: {exc}"}
        self._session_metadata.clear()
        try:
            self._metadata_path.unlink(missing_ok=True)
        except OSError as exc:
            return {"success": False, "error": f"Could not reset local run metadata: {exc}"}
        self.session_id = None
        self.source_session_id = None
        self.input_path = None
        self.run_dir = None
        return {"success": True, "files_preserved": True}

    def clear_all_sessions(self) -> dict[str, Any]:
        """Remove every completed run and reset the session auto-increment."""
        with self._lock:
            if self.process and self.process.poll() is None:
                return {"success": False, "error": "Stop the active run before clearing history."}

        run_dirs: set[Path] = set()
        with self._connect() as conn:
            sessions = conn.execute("SELECT id, input_file_path FROM sessions ORDER BY id").fetchall()
            for session in sessions:
                session_id = int(session["id"])
                metadata = self._session_metadata.get(str(session_id), {})
                candidates = [metadata.get("run_dir", "")]
                stored_input = str(session["input_file_path"] or "")
                if Path(stored_input).name.startswith(("input_source_", "input_restored_")):
                    candidates.append(str(Path(stored_input).parent))
                rows = conn.execute(
                    "SELECT download_folder, output_subdir FROM po_downloads WHERE session_id = ?",
                    (session_id,),
                ).fetchall()
                for row in rows:
                    if row["download_folder"]:
                        candidates.append(self._run_dir_from_download_folder(
                            str(row["download_folder"]), str(row["output_subdir"] or "")
                        ))
                roots = {self.default_download_root}
                if self.download_root:
                    roots.add(Path(self.download_root).expanduser())
                for root in roots:
                    report = root / f"report_session_{session_id}.xlsx"
                    if report.exists():
                        candidates.append(str(report.parent))
                    if root.exists():
                        candidates.extend(str(path) for path in root.glob(f"run_*/report_session_{session_id}.xlsx"))
                for candidate in candidates:
                    if not candidate:
                        continue
                    path = Path(candidate).expanduser()
                    if path.name.startswith("run_"):
                        run_dirs.add(path)

        deleted_dirs = 0
        for run_dir in sorted(run_dirs, key=lambda path: len(path.parts), reverse=True):
            try:
                resolved = run_dir.resolve()
                if resolved.name.startswith("run_") and resolved.parent != resolved and resolved.exists():
                    if resolved.is_symlink():
                        resolved.unlink()
                    else:
                        shutil.rmtree(resolved)
                    deleted_dirs += 1
            except OSError as exc:
                return {"success": False, "error": f"Could not remove run files: {exc}"}

        try:
            with self._connect() as conn:
                conn.execute("DELETE FROM retry_events")
                conn.execute("DELETE FROM retry_attempts")
                conn.execute("DELETE FROM po_downloads")
                conn.execute("DELETE FROM sessions")
                try:
                    conn.execute("DELETE FROM sqlite_sequence WHERE name IN ('sessions', 'po_downloads', 'retry_events', 'retry_attempts')")
                except sqlite3.OperationalError:
                    pass
                conn.commit()
        except sqlite3.Error as exc:
            return {"success": False, "error": f"Could not clear run history: {exc}"}

        self._session_metadata.clear()
        try:
            self._metadata_path.unlink(missing_ok=True)
        except OSError:
            pass
        return {"success": True, "deleted_dirs": deleted_dirs, "next_session_id": 1}

    def delete_session(self, session_id: int) -> dict[str, Any]:
        """Delete a completed run and its run-folder artifacts.

        The original input file is intentionally preserved when it lives
        outside the run folder; the archived copy inside the run folder is
        removed with the other artifacts.
        """
        session_id = int(session_id)
        with self._lock:
            if self.session_id == session_id and self.process and self.process.poll() is None:
                return {"success": False, "error": "Stop the active run before deleting it."}

        context = self._session_context(session_id)
        run_dir_value = str(context.get("run_dir") or "").strip()
        run_dir = Path(run_dir_value).expanduser() if run_dir_value else None
        if run_dir:
            try:
                resolved = run_dir.resolve()
            except OSError:
                resolved = run_dir
            # A run folder is always named run_*; this guard prevents a bad
            # metadata value from allowing deletion of an arbitrary directory.
            if resolved.name.startswith("run_") and resolved.parent != resolved and resolved.exists():
                try:
                    if resolved.is_symlink():
                        resolved.unlink()
                    else:
                        shutil.rmtree(resolved)
                except OSError as exc:
                    return {"success": False, "error": f"Could not remove run files: {exc}"}
            elif resolved.exists():
                return {"success": False, "error": "The stored run folder is not a safe run_* directory."}

        try:
            with self._connect() as conn:
                exists = conn.execute("SELECT 1 FROM sessions WHERE id = ?", (session_id,)).fetchone()
                if not exists:
                    return {"success": False, "error": "Run not found."}
                conn.execute("DELETE FROM retry_events WHERE session_id = ?", (session_id,))
                conn.execute("DELETE FROM retry_attempts WHERE session_id = ?", (session_id,))
                conn.execute("DELETE FROM po_attachments WHERE session_id = ?", (session_id,))
                conn.execute("DELETE FROM po_downloads WHERE session_id = ?", (session_id,))
                conn.execute("DELETE FROM sessions WHERE id = ?", (session_id,))
                conn.commit()
        except sqlite3.Error as exc:
            return {"success": False, "error": f"Could not remove run from history: {exc}"}

        self._session_metadata.pop(str(session_id), None)
        try:
            self._metadata_path.parent.mkdir(parents=True, exist_ok=True)
            self._metadata_path.write_text(json.dumps(self._session_metadata, indent=2), encoding="utf-8")
        except OSError:
            pass
        return {"success": True, "session_id": session_id, "run_dir": run_dir_value}

    def details(self, session_id: int) -> dict[str, Any]:
        with self._connect() as conn:
            session = conn.execute("SELECT * FROM sessions WHERE id = ?", (session_id,)).fetchone()
            rows = conn.execute("SELECT * FROM po_downloads WHERE session_id = ? ORDER BY id", (session_id,)).fetchall()
            retry_events = conn.execute("SELECT * FROM retry_events WHERE session_id = ? ORDER BY id DESC", (session_id,)).fetchall()
        session_data = dict(session) if session else None
        if session_data:
            session_data.pop("input_file_blob", None)
            metadata = self._session_context(session_id)
            session_data["input_file_path"] = session_data.get("input_file_path") or metadata.get("input_path", "")
            session_data["run_dir"] = str(metadata.get("run_dir") or "")
            stats = self._stats(session_id)
            session_data["status"] = self._derived_status(stats, session_data.get("status", "PENDING"))
            session_data.update({
                "total_pos": stats["total"],
                "success_count": stats["success"],
                "error_count": stats["errors"],
                "pending_count": stats["pending"],
            })
        supplier_by_po = self._supplier_by_po_from_input(
            str(session_data.get("input_file_path", "")) if session_data else ""
        )
        po_data = []
        for row in rows:
            item = dict(row)
            item["coupa_url"] = self._coupa_url(item.get("po_number", ""))
            item["supplier"] = supplier_by_po.get(str(item.get("po_number") or "").strip(), "")
            if not item["supplier"]:
                for input_po, supplier in supplier_by_po.items():
                    if self._po_values_equal(input_po, item.get("po_number", "")):
                        item["supplier"] = supplier
                        break
            po_data.append(item)
        return {"session": session_data, "pos": po_data, "retry_events": [dict(event) for event in retry_events]}

    def export_report(self, session_id: int, destination: Optional[str] = None) -> dict[str, Any]:
        details = self.details(session_id)
        if not details["session"]:
            return {"success": False, "error": "Session not found."}
        context = self._session_context(int(session_id))
        existing = None
        run_dir = str(context.get("run_dir") or self.run_dir or "").strip()
        if run_dir:
            candidate = Path(run_dir) / f"report_session_{session_id}.xlsx"
            if candidate.exists():
                existing = candidate
        if existing is None:
            root = Path(self.download_root or self.default_download_root)
            matches = list(root.glob(f"run_*/report_session_{session_id}.xlsx"))
            if matches:
                existing = matches[0]
        if existing:
            return {"success": True, "filepath": str(existing), "existing": True}
        default_output = Path(run_dir) / f"report_session_{session_id}.xlsx" if run_dir else Path.home() / "Downloads" / f"report_session_{session_id}.xlsx"
        output = Path(destination or default_output)
        output.parent.mkdir(parents=True, exist_ok=True)
        import pandas as pd
        pd.DataFrame(details["pos"]).to_excel(output, index=False)
        return {"success": True, "filepath": str(output)}
